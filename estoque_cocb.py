"""Estoque COCB — armazenamento em repositório privado GitHub.
Secrets: GITHUB_TOKEN, DATA_REPOSITORY, opcional DATA_BRANCH.
A cada gravação usa SHA para impedir sobrescrita de versões concorrentes.
"""
import base64
import json
import re
from datetime import datetime, timezone
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from urllib.parse import quote

from email_socorro import ArquivosEmail, destinatarios, CONFIG

import pandas as pd
import streamlit as st


central = st.query_params.get("pagina", "") == "central"
from pathlib import Path
from io import BytesIO
from PIL import Image
logo_10sul = Image.open(BytesIO(base64.b64decode((Path(__file__).resolve().parent / "assets" / "logo_10sul_base64.txt").read_text().strip().split(",")[-1])))
st.set_page_config(page_title="10 Sul" if central else "10 Sul | Estoque", page_icon=logo_10sul, layout="wide")
if central:
    st.markdown("""<style>
    .block-container {max-width:760px; padding-top:2rem;}
    [data-testid="stLinkButton"] a {min-height:70px; font-size:22px; border-radius:16px;}
    </style>""", unsafe_allow_html=True)
    st.image(logo_10sul, width=180)
    st.title("10 SUL")
    st.caption("Desenvolvido por Evandro Junior")
    st.subheader("Central de sistemas")
    st.link_button("📦 Estoque Socorro", "https://10sul-estoque-cocb.streamlit.app/?pagina=painel", use_container_width=True)
    st.link_button("🔧 Monitor Oficina", "https://10sul-monitor-gerencial-tah7ewvdikufrddrrysy56.streamlit.app/", use_container_width=True)
    st.info("Para ter um ícone no celular, adicione esta página à tela inicial com o nome 10 SUL.")
    st.stop()


UNIDADES = ("ARA", "MUC", "COCB", "NAM")

MATERIAIS = [
    ("27179902", "Válvula relê Facchini"),
    ("27174139", "Parafuso Facchini"),
    ("27174074", "Porca Facchini"),
    ("27130048", "Parafuso 1/1.8"),
    ("27114780", "Porca de 1.1/8"),
    ("27145672", "Mangotes"),
    ("27094293", "Válvula prévia"),
    ("27258592", "Bolsa de ar Manos"),
    ("27033363", "Bolsa de ar Sergomel"),
    ("27251038", "Arruela cônica"),
    ("27243269", "Colar de alinhamento"),
]

class ErroEstoque(Exception):
    pass

class BaseGitHub:
    def __init__(self):
        try:
            self.repo = str(st.secrets["DATA_REPOSITORY"]).strip()
            self.token = str(st.secrets["GITHUB_TOKEN"]).strip()
            self.branch = str(st.secrets.get("DATA_BRANCH", "main")).strip()
        except (KeyError, FileNotFoundError):
            raise ErroEstoque("Configure DATA_REPOSITORY e GITHUB_TOKEN nos Secrets.")
        if not re.fullmatch(r"[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+", self.repo) or not self.token:
            raise ErroEstoque("Repositório ou token inválido.")
        self.api = "https://api.github.com/repos/" + self.repo
        self.path = "/contents/dados/estoque_cocb.json"

    def requisitar(self, endpoint, metodo="GET", payload=None):
        headers = {"Authorization": "Bearer " + self.token,
                   "Accept": "application/vnd.github+json",
                   "X-GitHub-Api-Version": "2022-11-28",
                   "User-Agent": "10Sul-Estoque-COCB"}
        corpo = json.dumps(payload).encode("utf-8") if payload is not None else None
        if corpo is not None:
            headers["Content-Type"] = "application/json"
        try:
            with urlopen(Request(self.api + endpoint, data=corpo, headers=headers, method=metodo), timeout=25) as resp:
                return json.load(resp)
        except HTTPError as exc:
            if exc.code == 404 and metodo == "GET" and endpoint.startswith("/contents/"):
                return None
            if exc.code in (409, 422):
                raise ErroEstoque("Outra pessoa atualizou o estoque. Recarregue e repita a operação.") from None
            raise ErroEstoque(f"Falha de comunicação com o GitHub (HTTP {exc.code}).") from None
        except (OSError, URLError):
            raise ErroEstoque("Não foi possível acessar a base. Tente novamente.") from None

    def verificar_privado(self):
        info = self.requisitar("")
        if not info or not info.get("private"):
            raise ErroEstoque("A base deve estar em um repositório PRIVADO.")

    def ler(self):
        arquivo = self.requisitar(self.path + "?ref=" + quote(self.branch, safe=""))
        if arquivo is None:
            dados = {"schema": 1, "itens": {}, "movimentos": []}
            for ni, descricao in MATERIAIS:
                dados["itens"][ni] = {"descricao": descricao, "estoque_inicial": None, "saldo": 0, "minimo": None, "maximo": None}
            return dados, None
        try:
            dados = json.loads(base64.b64decode(arquivo["content"]))
            if dados.get("schema") != 1 or not isinstance(dados.get("itens"), dict) or not isinstance(dados.get("movimentos"), list):
                raise ValueError()
            normalizar(dados)
            return dados, arquivo["sha"]
        except (ValueError, KeyError, TypeError):
            raise ErroEstoque("Base de estoque inválida. Nenhuma informação foi modificada.") from None

    def salvar(self, dados, sha):
        conteudo = base64.b64encode(json.dumps(dados, ensure_ascii=False, indent=2).encode()).decode()
        payload = {"message": "Atualizar estoque COCB", "branch": self.branch, "content": conteudo}
        if sha:
            payload["sha"] = sha
        self.requisitar(self.path, "PUT", payload)


def normalizar(dados):
    """Preserva o saldo legado em COCB; as demais unidades começam vazias."""
    for item in dados["itens"].values():
        if "unidades" in item and "NAN" in item["unidades"] and "NAM" not in item["unidades"]:
            item["unidades"]["NAM"] = item["unidades"].pop("NAN")
        if "unidades" not in item:
            item["unidades"] = {
                u: {"saldo": item.get("saldo", 0) if u == "COCB" else 0,
                    "estoque_inicial": item.get("estoque_inicial") if u == "COCB" else None,
                    "minimo": item.get("minimo") if u == "COCB" else None,
                    "maximo": item.get("maximo") if u == "COCB" else None}
                for u in UNIDADES
            }
        for campo in ("minimo", "maximo"):
            geral = item["unidades"].get("COCB", {}).get(campo)
            if geral is None:
                geral = item.get(campo)
            if geral is None:
                geral = next((item["unidades"][u].get(campo) for u in UNIDADES if item["unidades"][u].get(campo) is not None), None)
            item[campo] = geral
            for u in UNIDADES:
                item["unidades"][u][campo] = geral
    for movimento in dados.get("movimentos", []):
        if movimento.get("unidade") == "NAN":
            movimento["unidade"] = "NAM"

def situacao(item):
    minimo, maximo = item.get("minimo"), item.get("maximo")
    saldo = item["saldo"]
    if minimo is not None and saldo < minimo:
        return "🔴 Abaixo do mínimo"
    if maximo is not None and saldo > maximo:
        return "🟠 Acima do máximo"
    if minimo is None or maximo is None:
        return "⚪ Não configurado"
    return "🟢 Dentro dos limites"

def itens_tabela(dados):
    normalizar(dados)
    linhas = []
    for ni, item in sorted(dados["itens"].items(), key=lambda x: x[1]["descricao"]):
        registro = item["unidades"][unidade]
        linhas.append({"NI": ni, "DESCRIÇÃO": item["descricao"],
                       **{u: item["unidades"][u]["saldo"] for u in UNIDADES},
                       "TOTAL": sum(item["unidades"][u]["saldo"] for u in UNIDADES),
                       "UNIDADE": unidade, "SALDO ATUAL": registro["saldo"],
                       "ESTOQUE INICIAL": registro["estoque_inicial"],
                       "MÍNIMO": registro.get("minimo"), "MÁXIMO": registro.get("maximo"),
                       "REPOR ATÉ MÁXIMO": None if registro.get("maximo") is None else max(0, registro["maximo"] - registro["saldo"]),
                       "SITUAÇÃO": situacao(registro)})
    return pd.DataFrame(linhas, columns=["NI", "DESCRIÇÃO", "MÍNIMO", "MÁXIMO", *UNIDADES, "TOTAL", "UNIDADE",
                                        "SALDO ATUAL", "REPOR ATÉ MÁXIMO", "ESTOQUE INICIAL", "SITUAÇÃO"])

def destacar_estoque_atual(tabela):
    """Vermelho estritamente abaixo do mínimo, por unidade."""
    def pintar(linha):
        minimo = linha.get("MÍNIMO")
        cores = []
        for coluna, valor in linha.items():
            limite = minimo
            if coluna == "TOTAL" and pd.notna(minimo):
                limite = minimo * len(UNIDADES)
            atual = coluna in (*UNIDADES, "TOTAL", "SALDO ATUAL", "ATUAL")
            cores.append("color: #b42318; font-weight: bold" if atual and pd.notna(limite) and pd.notna(valor) and valor < limite else "")
        return cores
    return tabela.style.apply(pintar, axis=1)


def alterar_ni(dados, ni_atual, novo_ni):
    novo_ni = str(novo_ni).strip()
    if not novo_ni:
        raise ErroEstoque("Informe o NI do material.")
    if novo_ni != ni_atual and novo_ni in dados["itens"]:
        raise ErroEstoque("Este NI já pertence a outro material.")
    if novo_ni != ni_atual:
        dados["itens"][novo_ni] = dados["itens"].pop(ni_atual)
        for movimento in dados["movimentos"]:
            if str(movimento.get("ni", "")) == ni_atual:
                movimento.setdefault("ni_original", ni_atual)
                movimento["ni"] = novo_ni
    return novo_ni


def registrar(dados, ni, tipo, quantidade, obs, unidade="COCB", go_carreta=""):
    normalizar(dados)
    material = dados["itens"].get(ni)
    if material is None:
        raise ErroEstoque("Material não encontrado.")
    if unidade not in UNIDADES or tipo not in ("ENTRADA", "SAÍDA"):
        raise ErroEstoque("Unidade ou movimento inválido.")
    go_carreta = str(go_carreta).strip().upper()
    if tipo == "SAÍDA" and not go_carreta:
        raise ErroEstoque("Informe o GO da carreta que receberá o material.")
    item = material["unidades"][unidade]
    if quantidade <= 0:
        raise ErroEstoque("Informe uma quantidade positiva.")
    if tipo == "SAÍDA" and item["estoque_inicial"] is None:
        raise ErroEstoque("Registre a primeira ENTRADA nesta unidade antes de realizar saídas.")
    saldo = item["saldo"] + (quantidade if tipo == "ENTRADA" else -quantidade)
    if saldo < 0:
        raise ErroEstoque(f"Saldo insuficiente. Disponível: {item['saldo']}.")
    primeira = tipo == "ENTRADA" and item["estoque_inicial"] is None
    if primeira:
        item["estoque_inicial"] = quantidade
    item["saldo"] = saldo
    if unidade == "COCB":
        material["saldo"] = saldo
        material["estoque_inicial"] = item["estoque_inicial"]
    dados["movimentos"].append({
        "data_hora": datetime.now(timezone.utc).isoformat(),
        "ni": ni, "descricao": material["descricao"], "unidade": unidade, "movimento": tipo,
        "quantidade": quantidade, "saldo_apos": saldo,
        "primeira_entrada": primeira, "go_carreta": go_carreta if tipo == "SAÍDA" else "", "observacao": obs.strip()})
    return primeira, saldo



def registrar_inventario(dados, ni, quantidade, responsavel, obs, unidade="COCB"):
    normalizar(dados)
    material = dados["itens"].get(ni)
    if material is None or unidade not in UNIDADES:
        raise ErroEstoque("Material ou unidade inválido.")
    if isinstance(quantidade, bool) or not isinstance(quantidade, int) or quantidade < 0:
        raise ErroEstoque("Informe uma quantidade física inteira, igual ou maior que zero.")
    responsavel = str(responsavel).strip()
    if not responsavel:
        raise ErroEstoque("Informe o responsável pela conferência.")
    item = material["unidades"][unidade]
    anterior = item["saldo"]
    diferenca = quantidade - anterior
    item["saldo"] = quantidade
    if item["estoque_inicial"] is None:
        item["estoque_inicial"] = quantidade
    if unidade == "COCB":
        material["saldo"] = quantidade
        material["estoque_inicial"] = item["estoque_inicial"]
    dados["movimentos"].append({
        "data_hora": datetime.now(timezone.utc).isoformat(),
        "ni": ni, "descricao": material["descricao"], "unidade": unidade,
        "movimento": "INVENTÁRIO", "quantidade": abs(diferenca),
        "saldo_anterior": anterior, "quantidade_fisica": quantidade,
        "diferenca": diferenca, "saldo_apos": quantidade,
        "responsavel": responsavel, "observacao": str(obs).strip(),
        "primeira_entrada": False, "go_carreta": ""})
    return anterior, diferenca


def registrar_inventario_lote(dados, linhas, responsavel, observacao="", unidade="COCB"):
    """Confere vários itens sem aplicar ajustes parciais."""
    import copy
    import math
    import uuid
    preparados = []
    usados = set()
    for linha in linhas:
        valor = linha.get("QUANTIDADE FÍSICA")
        if valor is None or pd.isna(valor):
            continue
        if isinstance(valor, bool):
            raise ErroEstoque("Informe quantidades físicas inteiras.")
        try:
            quantidade = float(valor)
        except (ValueError, TypeError):
            raise ErroEstoque("Informe quantidades físicas inteiras.") from None
        if not math.isfinite(quantidade) or quantidade < 0 or quantidade != int(quantidade):
            raise ErroEstoque("A quantidade física deve ser inteira, igual ou maior que zero.")
        ni = str(linha.get("NI", "")).strip()
        if ni in usados:
            raise ErroEstoque("O mesmo material foi informado mais de uma vez.")
        usados.add(ni)
        obs = str(linha.get("OBSERVAÇÃO") or "").strip() or str(observacao).strip()
        preparados.append((ni, int(quantidade), obs))
    if not preparados:
        raise ErroEstoque("Preencha a quantidade física de pelo menos um material.")
    novos = copy.deepcopy(dados)
    lote_id = str(uuid.uuid4())
    alterados = 0
    for ni, quantidade, obs in preparados:
        _, diferenca = registrar_inventario(novos, ni, quantidade, responsavel, obs, unidade)
        novos["movimentos"][-1]["lote_id"] = lote_id
        alterados += int(diferenca != 0)
    return novos, len(preparados), alterados


def exportar_tabela_excel(tabela, titulo):
    """Gera o arquivo Excel no servidor para download pelo aplicativo."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.sheet_view.showGridLines = False
    ws.append([str(c) for c in tabela.columns])
    for valores in tabela.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v.item() if hasattr(v, "item") else v for v in valores])
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="15364B")
        celula.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    linha = Side(style="thin", color="D6DEE6")
    for linha_dados in ws.iter_rows(min_row=2):
        for celula in linha_dados:
            celula.alignment = Alignment(vertical="center", wrap_text=True)
            celula.border = Border(bottom=linha)
            if celula.row % 2 == 0:
                celula.fill = PatternFill("solid", fgColor="F0F4F7")
            if isinstance(celula.value, (int, float)):
                celula.number_format = "#,##0"
    if "MÍNIMO" in tabela.columns:
        for numero, (_, item) in enumerate(tabela.iterrows(), start=2):
            minimo = item["MÍNIMO"]
            for indice, nome in enumerate(tabela.columns, start=1):
                limite = minimo * len(UNIDADES) if nome == "TOTAL" and pd.notna(minimo) else minimo
                if nome in (*UNIDADES, "TOTAL", "SALDO ATUAL", "ATUAL") and pd.notna(limite) and pd.notna(item[nome]) and item[nome] < limite:
                    ws.cell(numero, indice).font = Font(bold=True, color="B42318")
    for indice, nome in enumerate(tabela.columns, 1):
        tamanho = max([len(str(nome)), *[len(str(v)) for v in tabela[nome].dropna()]])
        ws.column_dimensions[get_column_letter(indice)].width = min(55, max(15, tamanho + 3))
    ws.freeze_panes = "C2"
    if len(tabela.columns):
        ws.auto_filter.ref = ws.dimensions
    arquivo = BytesIO()
    wb.save(arquivo)
    return arquivo.getvalue()


def exportar_criticos_excel(tabela, unidade):
    from io import BytesIO
    from zoneinfo import ZoneInfo
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Itens críticos"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:G1")
    ws["A1"] = "10 SUL — ITENS ABAIXO DO MÍNIMO"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="15364B")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:G2")
    ws["A2"] = "Unidade: " + unidade + " | " + datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M")
    cabecalhos = ["NI", "DESCRIÇÃO", "MÍNIMO", "MÁXIMO", "ESTOQUE ATUAL", "REPOR ATÉ MÁXIMO", "UNIDADE"]
    ws.append([])
    ws.append(cabecalhos)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="15364B")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 32
    for numero, (_, item) in enumerate(tabela.iterrows(), start=5):
        atual, maximo = int(item["SALDO ATUAL"]), item["MÁXIMO"]
        ws.append([str(item["NI"]), str(item["DESCRIÇÃO"]), item["MÍNIMO"],
                   None if pd.isna(maximo) else maximo, atual,
                   None if pd.isna(maximo) else max(0, maximo - atual), unidade])
        for c in ws[numero]:
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if numero % 2:
                c.fill = PatternFill("solid", fgColor="F0F4F7")
            if c.column in (3, 4, 5, 6):
                c.number_format = "#,##0"
        if pd.notna(item["MÍNIMO"]) and atual < item["MÍNIMO"]:
            ws.cell(numero, 5).font = Font(bold=True, color="B42318")
        ws.cell(numero, 6).font = Font(bold=True, color="B76E00")
        ws.row_dimensions[numero].height = 30
    for coluna, largura in enumerate((16, 48, 13, 13, 19, 23, 13), start=1):
        ws.column_dimensions[get_column_letter(coluna)].width = largura
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:G{ws.max_row}"
    ws.print_title_rows = "1:4"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:G{ws.max_row}"
    arquivo = BytesIO()
    wb.save(arquivo)
    return arquivo.getvalue()


def gerar_imagem_reposicao(dados, unidade):
    from io import BytesIO
    import textwrap
    from zoneinfo import ZoneInfo
    from PIL import Image, ImageDraw, ImageFont

    normalizar(dados)
    unidades_relatorio = ("MUC", "COCB")
    grupos = []
    for campo, titulo, cor in (
        ("maximo", "ITENS ABAIXO DO MÁXIMO", "#b76e00"),
    ):
        linhas = []
        for ni, material in sorted(dados["itens"].items(), key=lambda x: x[1]["descricao"]):
            reposicao = []
            abaixo = False
            for u in unidades_relatorio:
                estoque = material["unidades"][u]
                limite = estoque.get(campo)
                critico = limite is not None and estoque["saldo"] < limite
                abaixo = abaixo or critico
                maximo = estoque.get("maximo")
                reposicao.append(max(0, limite - estoque["saldo"]) if limite is not None else None)
            if abaixo:
                descricao = textwrap.wrap(material["descricao"], width=28) or [""]
                linhas.append((ni, descricao, reposicao, [material["unidades"][u] for u in unidades_relatorio]))
        grupos.append((titulo, cor, linhas))

    def fonte(tamanho):
        from pathlib import Path
        embutida = Path(__file__).resolve().parent / "assets" / "relatorio_font.b64"
        if embutida.exists():
            return ImageFont.truetype(BytesIO(base64.b64decode(embutida.read_text())), tamanho)
        for caminho in ("DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
            try:
                return ImageFont.truetype(caminho, tamanho)
            except OSError:
                pass
        return ImageFont.load_default(size=tamanho)

    altura = 400 + sum(225 + sum(max(48, len(l[1]) * 27 + 16) for l in linhas)
                        for _, _, linhas in grupos)
    imagem = Image.new("RGB", (1660, altura), "white")
    desenho = ImageDraw.Draw(imagem)
    desenho.rectangle((0, 0, 1660, 140), fill="#15364b")
    desenho.text((36, 25), "10 SUL | NECESSIDADE DE REPOSIÇÃO", font=fonte(32), fill="white")
    agora = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M")
    desenho.text((36, 82), f"Unidades: MUC e COCB  |  Atualizado em {agora}", font=fonte(23), fill="white")
    desenho.text((36, 155), "A REPOR = MÁXIMO − ATUAL (mínimo zero). Atual vermelho: abaixo do mínimo.", font=fonte(19), fill="#52616b")
    y = 200
    for titulo, cor, linhas in grupos:
        desenho.text((36, y), f"{titulo} ({len(linhas)})", font=fonte(26), fill=cor)
        y += 45
        desenho.rectangle((30, y, 1630, y + 65), fill="#eaf0f4")
        for x, texto in ((42, "NI"), (210, "MATERIAL"), (600, "MÍNIMO"), (710, "MÁXIMO"), (1500, "TOTAL"), (830, "ESTOQUE")):
            desenho.text((x, y + 8), texto, font=fonte(19), fill="#15364b")
        for x, nome in zip((1020, 1260), ("MUC", "COCB")):
            desenho.text((x + 40, y + 8), nome, font=fonte(21), fill="#15364b")
            desenho.text((x, y + 35), "ATUAL", font=fonte(17), fill="#15364b")
            desenho.text((x + 100, y + 35), "REPOR", font=fonte(17), fill="#15364b")
        desenho.text((830, y + 35), "ATUAL TOTAL", font=fonte(17), fill="#15364b")
        desenho.text((1500, y + 35), "REPOR", font=fonte(17), fill="#15364b")
        y += 65
        if not linhas:
            desenho.text((42, y + 12), "Nenhum item nesta condição.", font=fonte(21), fill="#52616b")
            y += 48
        totais = [0, 0]
        for indice, (ni, descricao, reposicao, estoques) in enumerate(linhas):
            h = max(48, len(descricao) * 27 + 16)
            desenho.rectangle((30, y, 1630, y + h), fill="#f5f7f9" if indice % 2 == 0 else "white")
            desenho.text((42, y + 10), ni, font=fonte(21), fill="#243746")
            for n, trecho in enumerate(descricao):
                desenho.text((210, y + 8 + n * 27), trecho, font=fonte(21), fill="#243746")
            for x, campo in ((600, "minimo"), (710, "maximo")):
                numero = estoques[0].get(campo)
                desenho.text((x, y + 10), "—" if numero is None else str(numero), font=fonte(22), fill="#243746")
            for x, estoque, repor in zip((1020, 1260), estoques, reposicao):
                desenho.text((x, y + 10), str(estoque["saldo"]), font=fonte(22), fill="#b42318" if estoque.get("minimo") is not None and estoque["saldo"] < estoque["minimo"] else "#243746")
                desenho.text((x + 100, y + 10), "—" if repor is None else str(repor), font=fonte(22), fill=cor if repor else "#52616b")
            desenho.text((830, y + 10), str(sum(estoque["saldo"] for estoque in estoques)), font=fonte(22), fill="#b42318" if all(e.get("minimo") is not None for e in estoques) and sum(e["saldo"] for e in estoques) < sum(e["minimo"] for e in estoques) else "#243746")
            desenho.text((1500, y + 10), str(sum(n or 0 for n in reposicao)), font=fonte(22), fill=cor)
            totais = [total + (numero or 0) for total, numero in zip(totais, reposicao)]
            desenho.line((30, y + h - 1, imagem.width - 30, y + h - 1), fill="#c6d2dc", width=2)
            y += h
        desenho.rectangle((30, y, 1630, y + 52), fill="#15364b")
        desenho.text((42, y + 14), "TOTAL A REPOR", font=fonte(21), fill="white")
        for x, numero in zip((1120, 1360, 1500), [*totais, sum(totais)]):
            desenho.text((x, y + 14), str(numero), font=fonte(22), fill="white")
        y += 77
    desenho.text((36, y), "Reposição calculada com base no estoque máximo de cada material.",
                 font=fonte(19), fill="#52616b")
    imagem = imagem.crop((0, 0, 1660, y + 50))
    arquivo = BytesIO()
    imagem.save(arquivo, format="PNG")
    return arquivo.getvalue()


def botao_compartilhar_imagem(png):
    import streamlit.components.v1 as components
    imagem_base64 = base64.b64encode(png).decode("ascii")
    html = """<!doctype html><html lang="pt-BR"><meta charset="utf-8">
<style>
body{margin:0;font-family:Arial,sans-serif}button{width:100%;padding:13px;border:0;border-radius:8px;background:#128c7e;color:white;font-size:16px;cursor:pointer}button:disabled{opacity:.65}p{font-size:13px;color:#52616b;margin:8px 0}
</style>
<button id="share">Compartilhar imagem no WhatsApp</button>
<p id="status" role="status">Escolha o WhatsApp e o contato na tela de compartilhamento.</p>
<script>
const bytes=Uint8Array.from(atob("__PNG__"),c=>c.charCodeAt(0));
const file=new File([bytes],"reposicao_10sul.png",{type:"image/png"});
const button=document.getElementById("share"), status=document.getElementById("status");
const navigators=[navigator];
try { if(window.parent!==window) navigators.unshift(window.parent.navigator); } catch(e) {}
button.onclick=async()=>{
 const sharing=navigators.find(n=>{try{return typeof n.share==="function"&&typeof n.canShare==="function"&&n.canShare({files:[file]});}catch(e){return false;}});
 if(!sharing){status.textContent="Este navegador não permite compartilhar imagens diretamente. Use Baixar imagem PNG e anexe no WhatsApp.";return;}
 button.disabled=true;
 try{
  await sharing.share({files:[file],title:"10 SUL — Necessidade de reposição"});
  status.textContent="Imagem compartilhada com o aplicativo escolhido.";
 }catch(e){
  status.textContent=e.name==="AbortError"?"Compartilhamento cancelado. Você pode tentar novamente.":"Não foi possível abrir o compartilhamento. Use Baixar imagem PNG e anexe no WhatsApp.";
 }finally{button.disabled=false;}
};
</script></html>""".replace("__PNG__", imagem_base64)
    components.html(html, height=120, scrolling=False)

def painel_email_socorro():
    arquivos = ArquivosEmail(base.requisitar, base.branch)
    try:
        cfg, cfg_sha = arquivos.configuracao()
    except ErroEstoque as exc:
        st.error(str(exc))
        return
    st.subheader("Enviar relatório por e-mail")
    st.caption("Relatório de MUC e COCB. Automático: terça e quinta às 9h (Brasília), mesmo com o aplicativo fechado. O processamento pode atrasar no GitHub.")
    with st.form("email_socorro"):
        emails = st.text_area("Destinatários", value="\n".join(cfg["destinatarios"]),
                              help="Separe por linha, vírgula ou ponto e vírgula. Você pode adicionar ou remover e-mails.")
        automatico = st.checkbox("Manter envio automático", value=cfg.get("automatico", True))
        salvar = st.form_submit_button("Salvar destinatários e agendamento")
        enviar = st.form_submit_button("Enviar agora por Resend", type="primary")
    if salvar or enviar:
        try:
            lista = destinatarios(emails)
            if salvar:
                arquivos.salvar(CONFIG, {"destinatarios": lista, "automatico": automatico,
                    "atualizado_em": datetime.now(timezone.utc).isoformat()}, cfg_sha)
                st.success("Destinatários e opção de envio automático salvos.")
            else:
                ident = arquivos.solicitar(lista)
                st.session_state["socorro_ultimo_envio"] = ident
                st.success("Envio solicitado. Aguarde o processamento e atualize o histórico para conferir o resultado.")
                st.caption("Protocolo: " + ident)
        except (ValueError, ErroEstoque) as exc:
            st.error(str(exc))
    st.subheader("Histórico de envios")
    st.button("Atualizar histórico", key="atualizar_historico_email")
    try:
        registros = arquivos.historico()
        pendente = st.session_state.get("socorro_ultimo_envio")
        if pendente and not any(x.get("id") == pendente for x in registros):
            st.info("Última solicitação na fila: " + pendente)
        if registros:
            tabela = pd.DataFrame([{
                "Data (Brasília)": pd.to_datetime(x["solicitado_em"], utc=True).tz_convert("America/Sao_Paulo").strftime("%d/%m/%Y %H:%M"),
                "Tipo": x.get("tipo", ""), "Destinatários": "; ".join(x.get("destinatarios", [])),
                "Status": {"aceito": "Aceito pelo Resend", "falha": "Falha", "incerto": "Sem confirmação", "processando": "Processando"}.get(x.get("status"), x.get("status", "")),
                "ID Resend": x.get("resend_id", ""), "Detalhes": x.get("erro", ""), "Protocolo": x.get("id", "")
            } for x in registros])
            st.dataframe(tabela, hide_index=True, use_container_width=True)
            st.caption("Aceito pelo Resend confirma o recebimento pelo provedor. A entrega na caixa postal deve ser consultada no Resend.")
        else:
            st.info("Nenhum envio registrado nesta versão. Os testes anteriores permanecem no GitHub Actions.")
    except ErroEstoque as exc:
        st.error("Não foi possível consultar o histórico: " + str(exc))


@st.dialog("Relatório de reposição", width="large")
def abrir_relatorio_reposicao(dados, unidade):
    png = gerar_imagem_reposicao(dados, unidade)
    st.image(png, use_container_width=True)
    botao_compartilhar_imagem(png)
    st.download_button("Baixar imagem PNG", png, "reposicao_todas_unidades.png", mime="image/png")
    st.caption("Para enviar por e-mail, use Envio por e-mail e histórico no painel.")


st.title("📦 10 SUL • CONTROLE DE ESTOQUE")
st.caption("Desenvolvido por Evandro Junior")
try:
    base = BaseGitHub()
    base.verificar_privado()
    dados, sha = base.ler()
except ErroEstoque as exc:
    st.error(str(exc))
    st.stop()

normalizar(dados)
unidade = st.selectbox("Unidade", UNIDADES, index=2)
somente_movimentacao = (st.query_params.get("acesso", "") == "movimentacao"
                         or st.query_params.get("pagina", "") == "movimentacao")
if not somente_movimentacao and st.button("📷 Relatório para enviar ao cliente", type="primary", use_container_width=True):
    abrir_relatorio_reposicao(dados, unidade)

pagina = st.query_params.get("pagina", "movimentacao" if somente_movimentacao else "painel")
if somente_movimentacao and pagina not in ("movimentacao", "inventario"):
    pagina = "movimentacao"
if pagina not in ("painel", "cadastro", "movimentacao", "inventario"):
    pagina = "painel"
if somente_movimentacao:
    a, b = st.columns(2)
    a.link_button("📦 Movimentação", "?acesso=movimentacao&pagina=movimentacao", use_container_width=True)
    b.link_button("🔎 Inventário", "?acesso=movimentacao&pagina=inventario", use_container_width=True)
else:
    a, b, c, d = st.columns(4)
    a.link_button("📊 Painel", "?pagina=painel", use_container_width=True)
    b.link_button("📋 Cadastro", "?pagina=cadastro", use_container_width=True)
    c.link_button("📦 Movimentação", "?pagina=movimentacao", use_container_width=True)
    d.link_button("🔎 Inventário", "?pagina=inventario", use_container_width=True)

if pagina == "painel":
    with st.expander("✉️ Envio por e-mail e histórico", expanded=False):
        painel_email_socorro()
    st.subheader("Painel de Estoque — " + unidade)
    tabela = itens_tabela(dados)
    criticos = tabela[tabela["SITUAÇÃO"] == "🔴 Abaixo do mínimo"].copy()
    configurados = sum(x["unidades"][unidade].get("minimo") is not None and x["unidades"][unidade].get("maximo") is not None for x in dados["itens"].values())
    a1, a2, a3 = st.columns(3)
    a1.metric("Materiais", len(dados["itens"]))
    a2.metric("Abaixo do mínimo", len(criticos))
    a3.metric("Limites definidos", configurados)
    st.subheader("🔴 Itens abaixo do mínimo")
    if criticos.empty:
        st.success("Nenhum material abaixo do mínimo.")
    else:
        criticos["REPOR ATÉ MÁXIMO"] = (criticos["MÁXIMO"] - criticos["SALDO ATUAL"]).clip(lower=0)
        st.dataframe(destacar_estoque_atual(criticos[["NI", "DESCRIÇÃO", "SALDO ATUAL", "MÍNIMO", "MÁXIMO", "REPOR ATÉ MÁXIMO"]]), hide_index=True, use_container_width=True)
        st.download_button("Exportar itens críticos em Excel", exportar_criticos_excel(criticos, unidade), f"estoque_critico_{unidade}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.subheader("Alertas de todas as unidades")
    alertas = []
    for ni, material in dados["itens"].items():
        for u in UNIDADES:
            estoque = material["unidades"][u]
            if situacao(estoque) == "🔴 Abaixo do mínimo":
                alertas.append({"NI": ni, "DESCRIÇÃO": material["descricao"], "UNIDADE": u,
                                "ATUAL": estoque["saldo"], "MÍNIMO": estoque["minimo"],
                                "MÁXIMO": estoque["maximo"]})
    if alertas:
        st.dataframe(destacar_estoque_atual(pd.DataFrame(alertas)), hide_index=True, use_container_width=True)
    else:
        st.info("Nenhum item abaixo do mínimo nas quatro unidades.")
    st.subheader("Consulta geral")
    apenas_criticos = st.checkbox("Mostrar apenas itens abaixo do mínimo")
    st.dataframe(destacar_estoque_atual(criticos if apenas_criticos else tabela), hide_index=True, use_container_width=True)
elif pagina == "cadastro":
    st.subheader("Cadastro de materiais")
    st.info("A primeira ENTRADA de cada material definirá o estoque inicial.")
    with st.form("cadastro", clear_on_submit=True):
        ni = st.text_input("NI (informado manualmente)")
        descricao = st.text_input("Descrição")
        minimo_novo = st.number_input("Estoque mínimo", min_value=0, step=1)
        maximo_novo = st.number_input("Estoque máximo", min_value=0, step=1)
        gravar = st.form_submit_button("Cadastrar", type="primary")
    if gravar:
        ni, descricao = ni.strip(), descricao.strip()
        if not ni or not descricao:
            st.error("Preencha NI e descrição.")
        elif ni in dados["itens"]:
            st.warning("Este NI já existe; nada foi alterado.")
        elif maximo_novo < minimo_novo:
            st.error("O máximo não pode ser menor que o mínimo.")
        else:
            dados["itens"][ni] = {"descricao": descricao, "estoque_inicial": None, "saldo": 0, "minimo": None, "maximo": None}
            normalizar(dados)
            dados["itens"][ni].update(minimo=int(minimo_novo), maximo=int(maximo_novo))
            for registro in dados["itens"][ni]["unidades"].values():
                registro.update(minimo=int(minimo_novo), maximo=int(maximo_novo))
            normalizar(dados)
            try:
                base.salvar(dados, sha)
                st.success("Material cadastrado.")
                st.rerun()
            except ErroEstoque as exc:
                st.error(str(exc))
    st.subheader("Editar item e limites gerais")
    if dados["itens"]:
        escolhas = {f"{ni} | {item['descricao']}": ni for ni, item in sorted(dados["itens"].items())}
        escolhido_limite = st.selectbox("Material para editar", list(escolhas))
        ni_editar = escolhas[escolhido_limite]
        material = dados["itens"][ni_editar]
        item_limite = material["unidades"][unidade]
        with st.form("limites_" + ni_editar + "_" + unidade):
            ni_editado = st.text_input("NI do item", value=ni_editar)
            descricao_editada = st.text_input("Descrição do item", value=material["descricao"])
            col_min, col_max = st.columns(2)
            minimo = col_min.number_input("Estoque mínimo", min_value=0, value=int(item_limite.get("minimo") or 0), step=1)
            maximo = col_max.number_input("Estoque máximo", min_value=0, value=int(item_limite.get("maximo") or 0), step=1)
            salvar_limites = st.form_submit_button("Salvar alterações", type="primary")
        if salvar_limites:
            ni_editado = ni_editado.strip()
            if not ni_editado:
                st.error("Informe o NI do material.")
            elif ni_editado != ni_editar and ni_editado in dados["itens"]:
                st.error("Este NI já pertence a outro material.")
            elif not descricao_editada.strip():
                st.error("Informe a descrição.")
            elif maximo < minimo:
                st.error("O estoque máximo não pode ser menor que o mínimo.")
            else:
                alterar_ni(dados, ni_editar, ni_editado)
                material["descricao"] = descricao_editada.strip()
                item_limite["minimo"] = int(minimo)
                item_limite["maximo"] = int(maximo)
                material.update(minimo=int(minimo), maximo=int(maximo))
                for registro in material["unidades"].values():
                    registro.update(minimo=int(minimo), maximo=int(maximo))
                normalizar(dados)
                try:
                    base.salvar(dados, sha)
                    st.rerun()
                except ErroEstoque as exc:
                    st.error(str(exc))
    st.subheader("Importar materiais")
    arquivo = st.file_uploader("Excel com NI e DESCRIÇÃO", type=["xlsx"])
    if arquivo:
        try:
            planilha = pd.read_excel(arquivo, dtype=str).fillna("")
            planilha.columns = [str(c).strip().upper().replace("Ç", "C").replace("Ã", "A") for c in planilha.columns]
            if not {"NI", "DESCRICAO"}.issubset(planilha.columns):
                st.error("A planilha precisa conter NI e DESCRIÇÃO.")
            else:
                st.dataframe(planilha[["NI", "DESCRICAO"]].head(30), hide_index=True)
                if st.button("Importar"):
                    novos = 0
                    for _, linha in planilha.iterrows():
                        ni, descricao = str(linha["NI"]).strip(), str(linha["DESCRICAO"]).strip()
                        if ni and descricao and ni not in dados["itens"]:
                            dados["itens"][ni] = {"descricao": descricao, "estoque_inicial": None, "saldo": 0, "minimo": None, "maximo": None}
                            novos += 1
                    if novos:
                        try:
                            base.salvar(dados, sha)
                            st.success(f"{novos} materiais importados.")
                            st.rerun()
                        except ErroEstoque as exc:
                            st.error(str(exc))
                    else:
                        st.info("Nenhum NI novo encontrado.")
        except Exception:
            st.error("Não foi possível ler o Excel.")
    st.subheader("Consulta de estoque")
    tabela = itens_tabela(dados)
    busca = st.text_input("Buscar NI ou descrição")
    if busca:
        tabela = tabela[tabela["NI"].str.contains(busca, case=False, regex=False) |
                        tabela["DESCRIÇÃO"].str.contains(busca, case=False, regex=False)]
    st.dataframe(tabela, use_container_width=True, hide_index=True)
    st.download_button("Exportar saldos em Excel", exportar_tabela_excel(tabela, "Saldos"), "saldos_estoque.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

elif pagina == "inventario":
    st.subheader("Inventário — " + unidade)
    if mensagem := st.session_state.pop("inventario_sucesso", None):
        st.success(mensagem)
    st.caption("Preencha a quantidade física dos materiais conferidos. Campo vazio mantém o saldo; zero informa que não há estoque físico.")
    tabela = itens_tabela(dados)
    if tabela.empty:
        st.info("Cadastre materiais antes de realizar o inventário.")
    else:
        chave = "inventario_tabela_" + unidade
        if chave not in st.session_state:
            linhas = tabela[["NI", "DESCRIÇÃO", "SALDO ATUAL"]].copy()
            linhas["QUANTIDADE FÍSICA"] = pd.Series(pd.NA, index=linhas.index, dtype="Int64")
            linhas["OBSERVAÇÃO"] = ""
            st.session_state[chave] = {"sha": sha, "linhas": linhas}
        referencia = st.session_state[chave]
        if st.button("Atualizar saldos e limpar conferência"):
            st.session_state.pop(chave, None)
            st.rerun()
        with st.form("inventario_lote_" + unidade + "_" + str(referencia["sha"])):
            responsavel = st.text_input("Responsável pela conferência")
            observacao = st.text_input("Observação geral (opcional)")
            preenchida = st.data_editor(
                referencia["linhas"], hide_index=True, use_container_width=True,
                disabled=["NI", "DESCRIÇÃO", "SALDO ATUAL"], num_rows="fixed",
                column_config={
                    "NI": st.column_config.TextColumn("NI", width="small"),
                    "DESCRIÇÃO": st.column_config.TextColumn("Material", width="large"),
                    "SALDO ATUAL": st.column_config.NumberColumn("Saldo no sistema", format="%d"),
                    "QUANTIDADE FÍSICA": st.column_config.NumberColumn("Quantidade física", min_value=0, step=1, format="%d"),
                    "OBSERVAÇÃO": st.column_config.TextColumn("Observação"),
                },
                key="editor_" + chave + "_" + str(referencia["sha"]),
                height=min(700, 40 + 35 * len(referencia["linhas"])))
            confirmar = st.form_submit_button("Salvar todos os itens conferidos", type="primary")
        if confirmar:
            try:
                if referencia["sha"] != sha:
                    raise ErroEstoque("O estoque mudou durante a conferência. Nenhum item foi ajustado. Atualize os saldos, confira novamente e repita o lançamento.")
                novos, conferidos, alterados = registrar_inventario_lote(
                    dados, preenchida.to_dict("records"), responsavel, observacao, unidade)
                base.salvar(novos, sha)
                st.session_state.pop(chave, None)
                st.session_state["inventario_sucesso"] = (
                    f"Inventário salvo: {conferidos} item(ns) conferido(s), {alterados} saldo(s) corrigido(s).")
                st.rerun()
            except ErroEstoque as exc:
                st.error(str(exc))
    st.subheader("Histórico de inventários")
    historico = [m for m in reversed(dados["movimentos"])
                 if m.get("movimento") == "INVENTÁRIO" and m.get("unidade", "COCB") == unidade]
    if historico:
        tabela_inventarios = pd.DataFrame(historico).reindex(columns=[
            "data_hora", "ni", "descricao", "unidade", "saldo_anterior",
            "quantidade_fisica", "diferenca", "saldo_apos", "responsavel", "observacao"])
        st.dataframe(tabela_inventarios, hide_index=True, use_container_width=True)
    else:
        st.info("Nenhum inventário registrado nesta unidade.")

elif pagina == "movimentacao":
    st.subheader("Movimentação")
    tabela = itens_tabela(dados)
    if tabela.empty:
        st.info("Cadastre materiais antes de movimentar.")
    else:
        opcoes = {f"{r['NI']} | {r['DESCRIÇÃO']} | Saldo: {r['SALDO ATUAL']}": r["NI"] for _, r in tabela.iterrows()}
        tipo = st.selectbox("Movimento", ["ENTRADA", "SAÍDA"], key="tipo_movimentacao")
        with st.form("movimentacao_" + tipo, clear_on_submit=True):
            escolhido = st.selectbox("Material", list(opcoes))
            quantidade = st.number_input("Quantidade", min_value=1, step=1)
            go_carreta = ""
            if tipo == "SAÍDA":
                go_carreta = st.text_input("GO da carreta (obrigatório)", placeholder="Ex.: 13795")
            observacao = st.text_input("Observação (opcional)")
            confirmar = st.form_submit_button("Registrar movimentação", type="primary")
        if confirmar:
            try:
                primeira, saldo = registrar(dados, opcoes[escolhido], tipo, int(quantidade), observacao, unidade, go_carreta)
                base.salvar(dados, sha)
                st.success(f"{'Estoque inicial registrado!' if primeira else 'Movimentação registrada!'} Saldo: {saldo}")
                st.rerun()
            except ErroEstoque as exc:
                st.error(str(exc))
    st.subheader("Histórico")
    movimentos = pd.DataFrame(reversed(dados["movimentos"]))
    if not movimentos.empty:
        movimentos["unidade"] = movimentos.get("unidade", pd.Series("COCB", index=movimentos.index)).fillna("COCB")
        movimentos["GO DA CARRETA"] = movimentos.get("go_carreta", pd.Series("", index=movimentos.index)).fillna("")
        movimentos = movimentos.drop(columns=["go_carreta"], errors="ignore")
        movimentos = movimentos[movimentos["unidade"] == unidade]
    st.dataframe(movimentos, use_container_width=True, hide_index=True)
    st.download_button("Exportar histórico em Excel", exportar_tabela_excel(movimentos, "Histórico"), "historico_estoque.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

